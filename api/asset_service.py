"""Equipment service: Excel bulk import and equipment-code lookup.

Manual registration (POST /api/assets) and Excel import (POST
/api/assets/import) share this module and the same ``Asset`` model — there is
no separate "imported equipment" entity.

Duplicate-code policy (deterministic, documented here and in
docs/downstream/equipment-maintenance.md):
- Equipment codes are matched case-insensitively on the trimmed value.
- Duplicate codes inside the uploaded file: the first occurrence is kept,
  later occurrences are rejected and reported.
- Codes that already exist in the database: rejected and reported.
- Existing records are never overwritten or merged by an import.
- A UNIQUE database index on ``asset.code`` (migration 27) backstops the
  application-level checks against concurrent writers.
"""

from dataclasses import dataclass, field
from io import BytesIO
from typing import Any, Dict, List, Optional, Set

from loguru import logger
from openpyxl import load_workbook

from open_notebook.database.repository import repo_query
from open_notebook.domain.asset import Asset
from open_notebook.exceptions import InvalidInputError, NotFoundError

REQUIRED_CODE_COLUMN = "Code"

# Excel header -> Asset field. Headers are normalized (trimmed, lowercased,
# internal whitespace collapsed) before lookup. Aliases cover only the exact
# column spellings from the equipment template (e.g. "ZONE-Description" vs
# "plant-description"); unknown columns are ignored, never guessed.
_EQUIPMENT_COLUMN_ALIASES: Dict[str, str] = {
    "code": "code",
    "main description": "name",
    "factory": "factory",
    "main function location": "location",
    "zone-description": "zone_description",
    "zone description": "zone_description",
    "site-description": "site_description",
    "site description": "site_description",
    "plant-description": "plant_description",
    "plant description": "plant_description",
    "main class": "main_class",
    "sub class": "sub_class",
    "type-description": "asset_type",
    "type description": "asset_type",
    "manufacture": "manufacturer",
    "manufacturer": "manufacturer",
    "model": "model",
}

_EQUIPMENT_FIELDS = (
    "code",
    "name",
    "factory",
    "location",
    "zone_description",
    "site_description",
    "plant_description",
    "main_class",
    "sub_class",
    "asset_type",
    "manufacturer",
    "model",
)


def normalize_equipment_code(code: Optional[str]) -> str:
    """Trim an equipment code; empty/blank becomes "" (validated downstream)."""
    return (code or "").strip()


def equipment_code_key(code: str) -> str:
    """Canonical duplicate-matching key: trimmed, case-insensitive."""
    return code.strip().upper()


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split()).lower()


def _cell_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    text = str(value).strip()
    return text if text else None


@dataclass
class EquipmentImportRow:
    row_number: int
    values: Dict[str, Optional[str]] = field(default_factory=dict)


@dataclass
class EquipmentImportIssue:
    row_number: int
    code: Optional[str]
    message: str


@dataclass
class EquipmentImportResult:
    total_rows: int
    valid_rows: List[EquipmentImportRow] = field(default_factory=list)
    issues: List[EquipmentImportIssue] = field(default_factory=list)
    imported_count: int = 0


def parse_equipment_workbook(content: bytes) -> EquipmentImportResult:
    """Parse and validate an equipment .xlsx file (no database writes).

    Raises InvalidInputError when the file is not a readable workbook or the
    required Code column is missing. Otherwise returns per-row validation:
    blank rows are skipped, rows without Code / Main Description and
    in-file duplicate codes become issues.
    """
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise InvalidInputError(
            "The uploaded file could not be read as an Excel (.xlsx) workbook."
        ) from e

    try:
        sheet = workbook.active
        if sheet is None:
            raise InvalidInputError("The Excel file contains no worksheet.")
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        raise InvalidInputError("The Excel file is empty.")

    header = [_normalize_header(cell) for cell in rows[0]]
    column_fields: List[Optional[str]] = [
        _EQUIPMENT_COLUMN_ALIASES.get(h) for h in header
    ]
    if "code" not in column_fields:
        raise InvalidInputError(
            "The Excel file must contain a 'Code' column with the equipment code."
        )

    result = EquipmentImportResult(total_rows=0)
    seen_codes: Set[str] = set()
    for index, raw_row in enumerate(rows[1:], start=2):
        values: Dict[str, Optional[str]] = {}
        for col, cell in enumerate(raw_row):
            if col >= len(column_fields):
                break
            target = column_fields[col]
            if target is None:
                continue
            text = _cell_text(cell)
            if text is not None:
                values[target] = text
        if not values:
            continue  # completely blank row
        result.total_rows += 1

        code = normalize_equipment_code(values.get("code"))
        name = (values.get("name") or "").strip()
        if not code:
            result.issues.append(
                EquipmentImportIssue(
                    row_number=index, code=None, message="Missing equipment Code."
                )
            )
            continue
        if not name:
            result.issues.append(
                EquipmentImportIssue(
                    row_number=index,
                    code=code,
                    message="Missing Main Description.",
                )
            )
            continue
        key = equipment_code_key(code)
        if key in seen_codes:
            result.issues.append(
                EquipmentImportIssue(
                    row_number=index,
                    code=code,
                    message=f"Duplicate code '{code}' inside the file; "
                    "only the first occurrence is kept.",
                )
            )
            continue
        seen_codes.add(key)
        values["code"] = code
        values["name"] = name
        result.valid_rows.append(EquipmentImportRow(row_number=index, values=values))
    return result


async def fetch_existing_code_keys() -> Set[str]:
    """Normalized equipment-code keys already present in the database."""
    try:
        rows = await repo_query("SELECT code FROM asset WHERE code != NONE")
    except Exception as e:
        logger.error(f"Error fetching existing equipment codes: {str(e)}")
        raise
    return {
        equipment_code_key(str(row["code"]))
        for row in rows
        if row.get("code") is not None
    }


def reject_existing_codes(
    result: EquipmentImportResult, existing_keys: Set[str]
) -> EquipmentImportResult:
    """Split preview rows against codes already in the database (no writes)."""
    normalized_existing = {equipment_code_key(code) for code in existing_keys}
    kept: List[EquipmentImportRow] = []
    for row in result.valid_rows:
        code = row.values.get("code") or ""
        if equipment_code_key(code) in normalized_existing:
            result.issues.append(
                EquipmentImportIssue(
                    row_number=row.row_number,
                    code=code,
                    message=f"Code '{code}' already exists and will not be "
                    "overwritten.",
                )
            )
        else:
            kept.append(row)
    result.valid_rows = kept
    return result


async def import_equipment_rows(
    rows: List[EquipmentImportRow],
) -> tuple[int, List[EquipmentImportIssue]]:
    """Persist validated rows one by one (row-level strategy, explicit).

    There is no multi-statement transaction in the repository layer (each
    repo_* call opens/closes its own connection), so rows are inserted
    individually. Every row was pre-validated; a write-time failure (e.g. a
    UNIQUE-index race with a concurrent writer) is caught per row, reported
    as an issue, and does not abort the remaining rows.
    """
    imported = 0
    issues: List[EquipmentImportIssue] = []
    for row in rows:
        try:
            asset = Asset(
                **{f: row.values.get(f) for f in _EQUIPMENT_FIELDS},
            )
            await asset.save()
            imported += 1
        except Exception as e:
            logger.warning(
                f"Skipping equipment row {row.row_number} "
                f"({row.values.get('code')}): {str(e)}"
            )
            issues.append(
                EquipmentImportIssue(
                    row_number=row.row_number,
                    code=row.values.get("code"),
                    message=f"Could not save this row: {str(e)}",
                )
            )
    return imported, issues


async def get_asset_by_code(code: str) -> Asset:
    """Fetch one equipment record by its code (case-insensitive, trimmed)."""
    normalized = normalize_equipment_code(code)
    if not normalized:
        raise InvalidInputError("Equipment code must be provided")
    rows = await repo_query(
        "SELECT * FROM asset WHERE string::uppercase(code OR '') == $code LIMIT 1",
        {"code": normalized.upper()},
    )
    if not rows:
        raise NotFoundError(f"Equipment with code '{normalized}' was not found")
    # Same construction as ObjectModel.get(): raw row straight into the model.
    return Asset(**rows[0])
