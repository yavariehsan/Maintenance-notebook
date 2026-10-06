"""Failure-mode service: Excel bulk import and failure-mode lookup.

Manual registration (POST /api/failure-modes) and Excel import (POST
/api/failure-modes/import) share this module and the same
``FailureMode`` model — there is no separate "imported failure mode"
entity. Every query here targets the dedicated ``failure_mode`` table
only, so failure-mode writes can never populate the equipment
``asset`` table (and equipment imports never touch this table).

Duplicate policy (mirrors the equipment registry):
- (code, label) pairs are matched case-insensitively on trimmed values.
- Duplicates inside the uploaded file: first occurrence kept, later
  ones rejected and reported.
- Pairs that already exist in the database: rejected and reported.
- Existing records are never overwritten or merged by an import.
- A UNIQUE database index on ``failure_mode.(code, label)``
  (migration 32) backstops the application checks against concurrent
  writers.
"""

from dataclasses import dataclass, field
from io import BytesIO
from typing import Dict, List, Optional, Set, Tuple

from loguru import logger
from openpyxl import load_workbook

from open_notebook.database.repository import repo_query
from open_notebook.domain.failure_mode import FailureMode
from open_notebook.exceptions import InvalidInputError

REQUIRED_CODE_COLUMN = "Code"
REQUIRED_LABEL_COLUMN = "Failure Mode"

# Excel header -> FailureMode field. Headers are normalized (trimmed,
# lowercased, internal whitespace collapsed) before lookup. Unknown
# columns are ignored, never guessed.
_FAILURE_MODE_COLUMN_ALIASES: Dict[str, str] = {
    "code": "code",
    "equipment code": "code",
    "failure mode": "label",
    "failure-mode": "label",
    "mode": "label",
    "label": "label",
    "description": "description",
}

_FAILURE_MODE_FIELDS = ("code", "label", "description")


def normalize_failure_mode_code(code: Optional[str]) -> str:
    """Trim an equipment code; empty/blank becomes "" (validated downstream)."""
    return (code or "").strip()


def failure_mode_key(code: str, label: str) -> Tuple[str, str]:
    """Canonical duplicate-matching key: trimmed, case-insensitive."""
    return (code.strip().upper(), label.strip().upper())


def _normalize_header(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split()).lower()


def _cell_text(value: object) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    text = str(value).strip()
    return text if text else None


@dataclass
class FailureModeImportRow:
    row_number: int
    values: Dict[str, Optional[str]] = field(default_factory=dict)


@dataclass
class FailureModeImportIssue:
    row_number: int
    code: Optional[str]
    message: str


@dataclass
class FailureModeImportResult:
    total_rows: int
    valid_rows: List[FailureModeImportRow] = field(default_factory=list)
    issues: List[FailureModeImportIssue] = field(default_factory=list)
    imported_count: int = 0


def parse_failure_mode_workbook(content: bytes) -> FailureModeImportResult:
    """Parse and validate a failure-mode .xlsx file (no database writes).

    Raises InvalidInputError when the file is not a readable workbook or
    a required column is missing. Otherwise returns per-row validation:
    blank rows are skipped, rows without Code / Failure Mode and in-file
    duplicate (code, label) pairs become issues.
    """
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise InvalidInputError(
            "The uploaded file could not be read as an Excel (.xlsx) workbook."
        ) from e

    try:
        sheet = workbook.active
        if sheet is None:
            raise InvalidInputError("The Excel file contains no worksheet.")
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        raise InvalidInputError("The Excel file is empty.")

    header = [_normalize_header(cell) for cell in rows[0]]
    column_fields: List[Optional[str]] = [
        _FAILURE_MODE_COLUMN_ALIASES.get(h) for h in header
    ]
    if "code" not in column_fields:
        raise InvalidInputError(
            "The Excel file must contain a 'Code' column with the equipment code."
        )
    if "label" not in column_fields:
        raise InvalidInputError(
            "The Excel file must contain a 'Failure Mode' column with the mode name."
        )

    result = FailureModeImportResult(total_rows=0)
    seen_keys: Set[Tuple[str, str]] = set()
    for index, raw_row in enumerate(rows[1:], start=2):
        values: Dict[str, Optional[str]] = {}
        for col, cell in enumerate(raw_row):
            if col >= len(column_fields):
                break
            target = column_fields[col]
            if target is None:
                continue
            text = _cell_text(cell)
            if text is not None:
                values[target] = text
        if not values:
            continue  # completely blank row
        result.total_rows += 1

        code = normalize_failure_mode_code(values.get("code"))
        label = (values.get("label") or "").strip()
        if not code:
            result.issues.append(
                FailureModeImportIssue(
                    row_number=index, code=None, message="Missing equipment Code."
                )
            )
            continue
        if not label:
            result.issues.append(
                FailureModeImportIssue(
                    row_number=index, code=code, message="Missing Failure Mode."
                )
            )
            continue
        key = failure_mode_key(code, label)
        if key in seen_keys:
            result.issues.append(
                FailureModeImportIssue(
                    row_number=index,
                    code=code,
                    message=f"Duplicate failure mode '{label}' for code "
                    f"'{code}' inside the file; only the first occurrence "
                    "is kept.",
                )
            )
            continue
        seen_keys.add(key)
        values["code"] = code
        values["label"] = label
        result.valid_rows.append(
            FailureModeImportRow(row_number=index, values=values)
        )
    return result


async def fetch_existing_mode_keys() -> Set[Tuple[str, str]]:
    """Normalized (code, label) pairs already present in the database."""
    try:
        rows = await repo_query(
            "SELECT code, label FROM failure_mode "
            "WHERE code != NONE AND label != NONE"
        )
    except Exception as e:
        logger.error(f"Error fetching existing failure modes: {str(e)}")
        raise
    return {
        failure_mode_key(str(row["code"]), str(row["label"]))
        for row in rows
        if row.get("code") is not None and row.get("label") is not None
    }


def reject_existing_modes(
    result: FailureModeImportResult, existing_keys: Set[Tuple[str, str]]
) -> FailureModeImportResult:
    """Split preview rows against pairs already in the database (no writes)."""
    normalized_existing = {
        (str(code).strip().upper(), str(label).strip().upper())
        for code, label in existing_keys
    }
    kept: List[FailureModeImportRow] = []
    for row in result.valid_rows:
        code = row.values.get("code") or ""
        label = row.values.get("label") or ""
        if failure_mode_key(code, label) in normalized_existing:
            result.issues.append(
                FailureModeImportIssue(
                    row_number=row.row_number,
                    code=code,
                    message=f"Failure mode '{label}' for code '{code}' "
                    "already exists and will not be overwritten.",
                )
            )
        else:
            kept.append(row)
    result.valid_rows = kept
    return result


async def import_failure_mode_rows(
    rows: List[FailureModeImportRow],
) -> tuple[int, List[FailureModeImportIssue]]:
    """Persist validated rows one by one into ``failure_mode`` only.

    There is no multi-statement transaction in the repository layer, so
    rows are inserted individually. A write-time failure (e.g. a UNIQUE
    race with a concurrent writer) is caught per row, reported as an
    issue, and does not abort the remaining rows.
    """
    imported = 0
    issues: List[FailureModeImportIssue] = []
    for row in rows:
        try:
            mode = FailureMode(
                **{f: row.values.get(f) for f in _FAILURE_MODE_FIELDS},
            )
            await mode.save()
            imported += 1
        except Exception as e:
            logger.warning(
                f"Skipping failure-mode row {row.row_number} "
                f"({row.values.get('code')}/{row.values.get('label')}): {str(e)}"
            )
            issues.append(
                FailureModeImportIssue(
                    row_number=row.row_number,
                    code=row.values.get("code"),
                    message=f"Could not save this row: {str(e)}",
                )
            )
    return imported, issues
