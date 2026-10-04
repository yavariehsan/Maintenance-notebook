"""Direct LLM insight flow (single call per scope): grouping, package,
prompt, validation, worker lifecycle, retrieval.

TDD RED step: none of the scope-insight functions exist yet.
"""
from io import BytesIO
from unittest.mock import AsyncMock, patch

import pytest

HEADERS = ["کد فرایندی", "شرح درخواست", "شرح تعمیر", "مکانیزم خرابی",
           "دلیل بروز عیب", "حالت خرابی",
           "تاریخ درخواست", "درخت تکنیکال"]

ROWS = [
    ["B138", "گیر کردن تعویض ابزار", "پاکت تعویض شد", "", "",
     "تعویض ابزار", "1402/09/01", "MACHINE TOOLS"],
    ["B138", "خطای تعویض ابزار", "سنسور تنظیم شد", "خرابی سنسور", "-",
     "تعویض ابزار", "1402/09/02", "MACHINE TOOLS"],
    ["B138", "توقف تعویض", "ریست شد", "", "علت نامشخص",
     "تعویض ابزار", "1402/09/03", "MACHINE TOOLS"],
    ["B138", "لرزش اسپیندل", "بلبرینگ تعویض شد", "", "",
     "خرابی اسپیندل", "1402/09/04", "MACHINE TOOLS"],
    ["C32", "گیر کردن ابزار", "پاکت تعویض شد", "", "",
     "تعویض ابزار", "1402/09/05", "MACHINE TOOLS"],
    ["", "", "", "", "", "", "", ""],
]


def _workbook_bytes() -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.append(HEADERS)
    for row in ROWS:
        sheet.append(row)
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def _headers_data():
    import openpyxl

    wb = openpyxl.load_workbook(BytesIO(_workbook_bytes()), read_only=True,
                                data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    return headers, [list(r) for r in rows[1:]]


def _guide(equipment="B138"):
    return "\n".join([
        "# ترتیب پیشنهادی تعمیرکار؛ نسخه عملیاتی",
        "### مرحله 1 — توقف و ثبت آلارم",
        "متن راهنما.",
        f"# دسته‌بندی کانون‌های اصلی خرابی {equipment}",
        "| کانون اصلی | سهم تاریخی |",
        "| --- | --- |",
        "| پاکت | 60 |",
        "راهنمای کاربردی تعمیر:",
        "| اولویت بررسی | اقدام پیشنهادی |",
        "| --- | --- |",
        "| 1 | بررسی پاکت |",
        "### موارد نامشخص / بدون علت قطعی",
        "علت نامشخص باقی ماند.",
    ])


# --- scope grouping ----------------------------------------------------------

def test_scope_groups_by_equipment_and_failure_mode():
    from api import llm_knowledge_service as svc

    headers, data = _headers_data()
    scopes = svc.plan_scope_insights("K", "Sheet1", headers, data)
    by_key = {(s["equipment"], s["failure_mode"]): s for s in scopes}
    assert set(by_key) == {
        ("B138", "تعویض ابزار"),
        ("B138", "خرابی اسپیندل"),
        ("C32", "تعویض ابزار"),
    }
    assert by_key[("B138", "تعویض ابزار")]["record_count"] == 3
    tool_ids = by_key[("B138", "تعویض ابزار")]["source_record_ids"]
    assert tool_ids == ["K-LLMROW-Sheet1-2", "K-LLMROW-Sheet1-3",
                        "K-LLMROW-Sheet1-4"]
    # No cross-equipment / cross-failure-mode mixing.
    assert by_key[("C32", "تعویض ابزار")]["source_record_ids"] == [
        "K-LLMROW-Sheet1-6"]
    assert by_key[("B138", "خرابی اسپیندل")]["source_record_ids"] == [
        "K-LLMROW-Sheet1-5"]


def test_scope_package_has_exactly_six_fields():
    from api import llm_knowledge_service as svc

    headers, data = _headers_data()
    scopes = svc.plan_scope_insights("K", "Sheet1", headers, data)
    scope = next(s for s in scopes if s["failure_mode"] == "تعویض ابزار"
                 and s["equipment"] == "B138")
    members = svc.plan_scope_records(scope, headers, data)
    assert len(members) == 3
    for member in members:
        assert set(member["fields"]) == set(svc.STAGE_A_FIELDS)
    text = svc.serialize_stage_a_record(members[0]["fields"])
    assert "تاریخ درخواست" not in text
    assert "MACHINE TOOLS" not in text


def test_scope_preserves_blanks_and_dashes():
    from api import llm_knowledge_service as svc

    headers, data = _headers_data()
    scopes = svc.plan_scope_insights("K", "Sheet1", headers, data)
    scope = next(s for s in scopes if s["record_count"] == 3)
    members = {m["source_record_id"]: m["fields"]
               for m in svc.plan_scope_records(scope, headers, data)}
    assert members["K-LLMROW-Sheet1-2"]["مکانیزم خرابی"] == ""
    assert members["K-LLMROW-Sheet1-3"]["دلیل بروز عیب"] == "-"


# --- prompt -----------------------------------------------------------------

def test_insight_prompt_has_required_sections_and_scope():
    from api import llm_knowledge_service as svc

    headers, data = _headers_data()
    scopes = svc.plan_scope_insights("K", "Sheet1", headers, data)
    scope = next(s for s in scopes if s["record_count"] == 3)
    members = svc.plan_scope_records(scope, headers, data)
    system, user = svc.build_insight_prompt(
        "B138", "تعویض ابزار",
        [{"source_record_id": m["source_record_id"], "fields": m["fields"]}
         for m in members])
    assert "# ترتیب پیشنهادی تعمیرکار" in system
    assert "# دسته‌بندی کانون‌های اصلی خرابی" in system
    assert "راهنمای کاربردی تعمیر:" in system
    assert "### موارد نامشخص / بدون علت قطعی" in system
    assert "B138" in user and "تعویض ابزار" in user
    assert "Record K-LLMROW-Sheet1-2:" in user
    assert "تاریخ درخواست" not in user
    assert "please provide" not in system.lower()


def test_insight_prompt_has_no_hardcoded_stats():
    from api import llm_knowledge_service as svc

    system, _ = svc.build_insight_prompt("B138", "تعویض ابزار", [])
    assert "60%" not in system and "٪۶۰" not in system
    assert "Do not fabricate percentages" in system or \
        "do not invent" in system.lower()


# --- validation -------------------------------------------------------------

def test_validate_insight_empty_fails():
    from api import llm_knowledge_service as svc

    for bad in ("", "   ", None):
        ok, problems, warnings = svc.validate_insight_markdown(bad)
        assert ok is False
        assert problems


def test_validate_insight_missing_sequence_fails():
    from api import llm_knowledge_service as svc

    ok, problems, _warnings = svc.validate_insight_markdown(
        "راهنمای کاربردی تعمیر:\nمتن بدون ترتیب.")
    assert ok is False
    assert any("sequence" in p for p in problems)


def test_validate_insight_missing_optional_sections_warns_only():
    from api import llm_knowledge_service as svc

    ok, problems, warnings = svc.validate_insight_markdown(
        "# ترتیب پیشنهادی تعمیرکار؛ نسخه عملیاتی\nمرحله 1.")
    assert ok is True
    assert problems == []
    assert warnings


def test_validate_insight_valid_returns_content_unchanged():
    from api import llm_knowledge_service as svc

    guide = _guide()
    ok, problems, _warnings = svc.validate_insight_markdown(guide)
    assert ok is True
    assert problems == []


# --- worker -----------------------------------------------------------------

def _build(status="queued"):
    return {
        "id": "llm_knowledge_build:1",
        "source_report_ids": ["repair_report:a"],
        "manifest": [{"report_id": "repair_report:a", "filename": "B.xlsx",
                      "analysis_key": "K"}],
        "status": status,
        "command_id": None,
        "model": "model:x",
        "prompt_version": "insight-v1",
        "error": None,
        "warnings": [],
        "record_count": None,
        "failed_record_count": None,
        "created": "2026-09-30T00:00:00",
        "started_at": None,
        "finished_at": None,
    }


def _report():
    return {"id": "repair_report:a", "filename": "B.xlsx", "analysis_key": "K"}


def _patches(worker, fake, saved, finished, build=None):
    worker.set_generate_fn(fake)
    return (
        patch.object(worker.llm_knowledge, "_get_build_internal",
                     new=AsyncMock(return_value=build or _build())),
        patch.object(worker.llm_knowledge, "mark_build_running", new=AsyncMock()),
        patch.object(worker.llm_knowledge, "save_stage_b_guide", new=saved),
        patch.object(worker.llm_knowledge, "_existing_stage_b_scopes",
                     new=AsyncMock(return_value=set())),
        patch.object(worker.llm_knowledge, "mark_build_finished",
                     new=AsyncMock(side_effect=lambda bid, **kw: finished.update(kw))),
        patch.object(worker.llm_knowledge, "mark_build_failed",
                     new=AsyncMock(side_effect=lambda bid, err: finished.update(
                         {"status": "failed", "error": err}))),
        patch.object(worker.llm_knowledge, "resolve_llm_model_id",
                     new=AsyncMock(return_value="model:x")),
        patch.object(worker.reports, "read_report_file",
                     new=AsyncMock(return_value=_workbook_bytes())),
        patch.object(worker.reports, "_get_report_internal",
                     new=AsyncMock(return_value=_report())),
    )


@pytest.mark.asyncio
async def test_worker_guides_every_scope_and_completes():
    from commands import llm_knowledge_commands as worker

    seen = {}

    async def _fake(system, user, scope_id, model_id, max_tokens):
        seen[scope_id] = (system, user, max_tokens)
        return _guide()

    saved, finished = [], {}

    async def _save(build_id, equipment, failure_mode, markdown, count,
                    batch_ids, source_ids, model, budget):
        saved.append({"equipment": equipment, "failure_mode": failure_mode,
                      "markdown": markdown, "count": count,
                      "source_ids": source_ids})
        return {"id": f"g:{equipment}"}

    try:
        patches = _patches(worker, _fake, _save, finished)
        with patches[0], patches[1], patches[2], patches[3], patches[4], \
                patches[5], patches[6], patches[7], patches[8]:
            result = await worker.generate_llm_knowledge_command(
                worker.GenerateLLMKnowledgeInput(build_id="llm_knowledge_build:1")
            )
    finally:
        worker.set_generate_fn(worker._default_batch_generate_fn)
    assert result.success is True
    assert len(saved) == 3  # three scopes, one LLM call each
    tool = next(s for s in saved if s["failure_mode"] == "تعویض ابزار"
                and s["equipment"] == "B138")
    assert tool["count"] == 3
    assert tool["markdown"] == _guide()  # stored AS-IS
    assert all(v[2] == 12000 for v in seen.values())  # stage_b budget
    assert finished["status"] == "completed"


@pytest.mark.asyncio
async def test_worker_partial_when_one_scope_fails_and_never_stuck():
    from commands import llm_knowledge_commands as worker

    async def _fake(system, user, scope_id, model_id, max_tokens):
        if "خرابی اسپیندل" in user:
            raise RuntimeError("provider boom")
        return _guide()

    saved, finished = [], {}

    async def _save(build_id, equipment, failure_mode, markdown, count,
                    batch_ids, source_ids, model, budget):
        saved.append(failure_mode)
        return {"id": "g:x"}

    try:
        patches = _patches(worker, _fake, _save, finished)
        with patches[0], patches[1], patches[2], patches[3], patches[4], \
                patches[5], patches[6], patches[7], patches[8]:
            result = await worker.generate_llm_knowledge_command(
                worker.GenerateLLMKnowledgeInput(build_id="llm_knowledge_build:1")
            )
    finally:
        worker.set_generate_fn(worker._default_batch_generate_fn)
    # Returns normally (no raise → never stuck running), partial terminal.
    assert result.success is True
    assert len(saved) == 2
    assert finished["status"] == "partial"
    assert any("خرابی اسپیندل" in w for w in finished.get("warnings", []))


@pytest.mark.asyncio
async def test_worker_all_scopes_fail_marks_failed_and_returns():
    from commands import llm_knowledge_commands as worker

    async def _fake(system, user, scope_id, model_id, max_tokens):
        raise RuntimeError("provider down")

    finished = {}

    async def _save(*args, **kwargs):
        raise AssertionError("must not persist failures")

    try:
        patches = _patches(worker, _fake, _save, finished)
        with patches[0], patches[1], patches[2], patches[3], patches[4], \
                patches[5], patches[6], patches[7], patches[8]:
            result = await worker.generate_llm_knowledge_command(
                worker.GenerateLLMKnowledgeInput(build_id="llm_knowledge_build:1")
            )
    finally:
        worker.set_generate_fn(worker._default_batch_generate_fn)
    assert result.success is False
    assert finished["status"] == "failed"
    assert finished.get("error")


@pytest.mark.asyncio
async def test_worker_unexpected_error_finalizes_instead_of_sticking():
    from commands import llm_knowledge_commands as worker

    async def _fake(system, user, scope_id, model_id, max_tokens):
        return _guide()

    finished = {}

    async def _save(*args, **kwargs):
        raise RuntimeError("database unavailable")

    try:
        patches = _patches(worker, _fake, _save, finished)
        with patches[0], patches[1], patches[2], patches[3], patches[4], \
                patches[5], patches[6], patches[7], patches[8]:
            result = await worker.generate_llm_knowledge_command(
                worker.GenerateLLMKnowledgeInput(build_id="llm_knowledge_build:1")
            )
    finally:
        worker.set_generate_fn(worker._default_batch_generate_fn)
    # No exception escapes → the build can never be left running.
    assert result.success is False
    assert finished["status"] == "failed"


@pytest.mark.asyncio
async def test_worker_resumes_completed_scopes():
    from commands import llm_knowledge_commands as worker

    calls = []

    async def _fake(system, user, scope_id, model_id, max_tokens):
        calls.append(scope_id)
        return _guide()

    saved, finished = [], {}

    async def _save(build_id, equipment, failure_mode, markdown, count,
                    batch_ids, source_ids, model, budget):
        saved.append((equipment, failure_mode))
        return {"id": "g:x"}

    try:
        patches = _patches(worker, _fake, _save, finished)
        done = {("B138", "تعویض ابزار")}
        with patches[0], patches[1], patches[2], \
                patch.object(worker.llm_knowledge, "_existing_stage_b_scopes",
                             new=AsyncMock(return_value=done)), \
                patches[4], patches[5], patches[6], patches[7], patches[8]:
            result = await worker.generate_llm_knowledge_command(
                worker.GenerateLLMKnowledgeInput(build_id="llm_knowledge_build:1")
            )
    finally:
        worker.set_generate_fn(worker._default_batch_generate_fn)
    assert result.success is True
    assert len(calls) == 2
    assert ("B138", "تعویض ابزار") not in saved
    assert finished["status"] == "completed"


# --- retrieval --------------------------------------------------------------

@pytest.mark.asyncio
async def test_assemble_guide_returns_scoped_guides_without_record_rows():
    from api import llm_knowledge_service as svc

    build = _build(status="completed")
    guides = [{
        "synthesis_id": "g:1",
        "build_id": "llm_knowledge_build:1",
        "equipment": "B138",
        "failure_mode": "تعویض ابزار",
        "guide_markdown": _guide(),
        "record_count": 3,
        "batch_ids": ["K-LLMSCOPE-ab12cd34"],
        "source_record_ids": ["K-LLMROW-Sheet1-2", "K-LLMROW-Sheet1-3",
                              "K-LLMROW-Sheet1-4", "OTHER-LLMROW-Sheet1-9"],
        "model": "model:x",
        "prompt_version": "insight-v1",
        "math_warnings": [],
    }]

    async def _repo(query, params=None):
        if "FROM llm_knowledge_build" in query:
            return [build]
        if "FROM llm_knowledge_record" in query:
            return []
        if "FROM llm_stage_b_guide" in query:
            return guides
        raise AssertionError(f"unexpected query: {query[:60]}")

    with patch.object(svc, "repo_query", new=_repo), \
        patch("api.repair_report_service._get_report_internal",
              new=AsyncMock(return_value=_report())):
        guide = await svc.assemble_llm_guide("llm_knowledge_build:1",
                                             "repair_report:a")
    assert guide["warnings"] == []
    assert len(guide["final_guides"]) == 1
    assert guide["final_guides"][0]["guide_markdown"] == _guide()
    assert guide["records"] == []


# --- persistence ------------------------------------------------------------

@pytest.mark.asyncio
async def test_save_stage_b_guide_builds_balanced_surrealql():
    """Regression: the CREATE statement must be syntactically valid.

    Production root cause: the query was assembled from an f-string head
    plus a plain-string tail, so the tail's ``}}`` reached SurrealDB as a
    literal double brace and every guide save failed with a parse error
    (``... time::now()}} RETURN AFTER``). Mocked repo_query in older
    tests never caught it because they never inspected the query text.
    """
    from api import llm_knowledge_service as svc

    seen = {}

    async def _repo(query, params=None):
        seen["query"] = query
        seen["params"] = params
        return [{"id": "llm_stage_b_guide:1"}]

    with patch.object(svc, "repo_query", new=_repo):
        await svc.save_stage_b_guide(
            "llm_knowledge_build:1", "B138", "تعویض ابزار", _guide(), 3,
            ["K-LLMSCOPE-ab12cd34"],
            ["K-LLMROW-Sheet1-2", "K-LLMROW-Sheet1-3", "K-LLMROW-Sheet1-4"],
            "model:x", 12000)

    query = seen["query"]
    assert "CREATE llm_stage_b_guide CONTENT" in query
    assert query.count("{") == query.count("}"), query[-60:]
    assert "}}" not in query
    assert query.rstrip().endswith("} RETURN AFTER")
    assert seen["params"]["val_record_count"] == 3
