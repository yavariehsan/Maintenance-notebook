"""Tests for the repair-report collection + analysis workflow (Milestone 6).

Covers, with mocked SurrealDB persistence (``repo_query`` seam):
upload validation, listing, detail, 10-row preview, analysis state
transitions, duplicate/concurrent protection, failure bookkeeping, and
no-leakage of storage paths.

Plus lifecycle consistency (Milestone 8): the report/run/command state
contract, submit-grace duplicate protection, worker-heartbeat liveness,
attach-failure cleanup, and repair-analysis rows on the Tasks endpoint.

Plus engine-level integration (real package pipeline, tmp files only):
multi-file aggregate determinism, record-ID namespacing, and the full
chain aggregate → engine → Troubleshooting DB → runtime adapter reads.

The real ``Sample-1.xlsx`` is never used here; only tiny synthetic
fixtures built in tmp directories.
"""

from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from unittest.mock import AsyncMock, patch

import pytest
from fastapi.testclient import TestClient

from api import repair_report_service as reports
from api.main import app

HEADERS = [
    "کد فرایندی",
    "تجهیز",
    "پیشوند درخواست",
    "شماره درخواست",
    "شرح درخواست",
    "شرح تعمیر",
    "حالت خرابی",
    "مکانیزم خرابی",
    "دلیل بروز عیب",
]


def _workbook_bytes(rows, headers=None):
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(headers or HEADERS)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _row(cells):
    full = list(cells) + [""] * (len(HEADERS) - len(cells))
    return full[: len(HEADERS)]


def _sample_rows(prefix="B", start=1, count=3, code="B104"):
    return [
        _row(
            [
                code,
                "mill",
                prefix,
                str(start + i),
                f"عیب شماره {start + i}",
                f"تعمیر انجام شد {start + i}. قطعه تعویض شد.",
                "روشن نشدن",
                "عدم بوت",
                "خرابی منبع تغذیه",
            ]
        )
        for i in range(count)
    ]


@pytest.fixture
def client():
    return TestClient(app)


def _report_row(**overrides):
    row = {
        "id": "repair_report:abc123",
        "filename": "cmms.xlsx",
        "stored_filename": "cmms.xlsx",
        "size_bytes": 1234,
        "sheet": "Sheet1",
        "column_count": 9,
        "data_rows": 3,
        "analysis_key": "a3f9c2e1",
        "analysis_state": "not_analyzed",
        "last_run_id": None,
        "last_completed_run_id": None,
        "last_error": None,
        "created": "2026-09-26T00:00:00",
        "updated": "2026-09-26T00:00:00",
    }
    row.update(overrides)
    return row


def _run_row(**overrides):
    row = {
        "id": "repair_analysis_run:run1",
        "report_ids": ["repair_report:abc123"],
        "manifest": [],
        "status": "queued",
        "command_id": None,
        "error": None,
        "record_count": None,
        "equipment_count": None,
        "failure_mode_count": None,
        "guide_count": None,
        "created": "2026-09-26T00:00:00",
        "started_at": None,
        "finished_at": None,
    }
    row.update(overrides)
    return row


def _repo_router(state):
    """Route repo_query calls to canned responses for router-level tests."""

    async def _repo(query, params=None):
        if query.startswith("CREATE repair_report"):
            return [state.get("created", _report_row())]
        if "FROM repair_report ORDER BY" in query:
            return state.get("reports", [_report_row()])
        if "FROM repair_report WHERE" in query:
            found = state.get("report")
            return [found] if found else []
        if query.startswith("CREATE repair_analysis_run"):
            return [_run_row()]
        if "FROM repair_analysis_run" in query:
            return state.get("runs", [])
        if "FROM command" in query:
            return state.get("commands", [])
        if query.startswith("UPDATE"):
            state.setdefault("updates", []).append((query, params))
            return []
        raise AssertionError(f"unexpected query: {query}")

    return _repo


def _iso_now_minus(minutes=0):
    return (datetime.now(timezone.utc) - timedelta(minutes=minutes)).isoformat()


def _live_run(**overrides):
    row = _run_row(status="processing", command_id="command:live")
    row.update(overrides)
    return row


# --- upload -----------------------------------------------------------------


@pytest.mark.asyncio
async def test_upload_registers_report(client, tmp_path, monkeypatch):
    monkeypatch.setattr(reports, "REPAIR_REPORTS_FOLDER", str(tmp_path))
    content = _workbook_bytes(_sample_rows(count=2))
    with patch.object(
        reports, "repo_query", new_callable=AsyncMock
    ) as mock_repo:
        mock_repo.side_effect = _repo_router(
            {"created": _report_row(data_rows=2, column_count=9)}
        )
        resp = client.post(
            "/api/repair-reports",
            files={
                "file": (
                    "cmms.xlsx",
                    content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert body["filename"] == "cmms.xlsx"
    assert body["analysis_state"] == "not_analyzed"
    assert body["data_rows"] == 2
    assert "stored_filename" not in body
    assert "repair-reports" not in resp.text
    assert "/data/" not in resp.text and "E:\\" not in resp.text


def test_upload_rejects_non_excel(client):
    resp = client.post(
        "/api/repair-reports",
        files={"file": ("notes.txt", b"hello", "text/plain")},
    )
    assert resp.status_code == 400


def test_upload_rejects_empty_file(client):
    resp = client.post(
        "/api/repair-reports",
        files={"file": ("empty.xlsx", b"", "application/octet-stream")},
    )
    assert resp.status_code == 400


def test_upload_rejects_unreadable_workbook(client):
    resp = client.post(
        "/api/repair-reports",
        files={"file": ("junk.xlsx", b"not a workbook", "application/octet-stream")},
    )
    assert resp.status_code == 400


def test_upload_rejects_headerless_workbook(client):
    import openpyxl

    book = openpyxl.Workbook()
    buffer = BytesIO()
    book.save(buffer)
    resp = client.post(
        "/api/repair-reports",
        files={"file": ("blank.xlsx", buffer.getvalue(), "application/octet-stream")},
    )
    assert resp.status_code == 400


# --- list / detail -----------------------------------------------------------


@pytest.mark.asyncio
async def test_list_reports(client):
    with patch.object(
        reports, "repo_query", new_callable=AsyncMock
    ) as mock_repo:
        mock_repo.side_effect = _repo_router(
            {"reports": [_report_row(), _report_row(id="repair_report:def456")]}
        )
        resp = client.get("/api/repair-reports")
    assert resp.status_code == 200
    assert [item["id"] for item in resp.json()] == [
        "repair_report:abc123",
        "repair_report:def456",
    ]


@pytest.mark.asyncio
async def test_get_report_detail_unknown(client):
    with patch.object(
        reports, "repo_query", new_callable=AsyncMock
    ) as mock_repo:
        mock_repo.side_effect = _repo_router({"report": None})
        resp = client.get("/api/repair-reports/repair_report:missing")
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_get_report_detail_includes_last_run(client):
    run = _run_row(status="completed")
    with patch.object(
        reports, "repo_query", new_callable=AsyncMock
    ) as mock_repo:
        mock_repo.side_effect = _repo_router(
            {
                "report": _report_row(
                    analysis_state="completed",
                    last_run_id="repair_analysis_run:run1",
                ),
                "runs": [run],
            }
        )
        # Detail looks up the run by id; route both run queries at it.
        resp = client.get("/api/repair-reports/repair_report:abc123")
    assert resp.status_code == 200
    body = resp.json()
    assert body["report"]["analysis_state"] == "completed"
    assert body["last_run"]["id"] == "repair_analysis_run:run1"
    assert body["last_run"]["status"] == "completed"


# --- preview ------------------------------------------------------------------


@pytest.mark.asyncio
async def test_preview_returns_ten_rows(tmp_path, monkeypatch, client):
    content = _workbook_bytes(_sample_rows(count=12))
    monkeypatch.setattr(reports, "REPAIR_REPORTS_FOLDER", str(tmp_path))
    stored = tmp_path / "cmms.xlsx"
    stored.write_bytes(content)
    with patch.object(
        reports, "repo_query", new_callable=AsyncMock
    ) as mock_repo:
        mock_repo.side_effect = _repo_router({"report": _report_row()})
        resp = client.get("/api/repair-reports/repair_report:abc123/preview")
    assert resp.status_code == 200
    body = resp.json()
    assert body["columns"][:4] == HEADERS[:4]
    assert len(body["rows"]) == 10
    assert body["total_data_rows"] == 12
    assert body["truncated"] is True
    # Persian text survives; empty cells are null-friendly.
    assert "عیب شماره 1" in body["rows"][0][4]


@pytest.mark.asyncio
async def test_preview_short_file_not_truncated(tmp_path, monkeypatch, client):
    content = _workbook_bytes(_sample_rows(count=3))
    monkeypatch.setattr(reports, "REPAIR_REPORTS_FOLDER", str(tmp_path))
    (tmp_path / "cmms.xlsx").write_bytes(content)
    with patch.object(
        reports, "repo_query", new_callable=AsyncMock
    ) as mock_repo:
        mock_repo.side_effect = _repo_router({"report": _report_row()})
        resp = client.get("/api/repair-reports/repair_report:abc123/preview")
    assert resp.status_code == 200
    body = resp.json()
    assert len(body["rows"]) == 3
    assert body["truncated"] is False


@pytest.mark.asyncio
async def test_preview_unknown_report(client):
    with patch.object(
        reports, "repo_query", new_callable=AsyncMock
    ) as mock_repo:
        mock_repo.side_effect = _repo_router({"report": None})
        resp = client.get("/api/repair-reports/repair_report:missing/preview")
    assert resp.status_code == 404


# --- analysis lifecycle --------------------------------------------------------


@pytest.mark.asyncio
async def test_start_analysis_empty_collection(client):
    with patch.object(
        reports, "repo_query", new_callable=AsyncMock
    ) as mock_repo:
        mock_repo.side_effect = _repo_router({"reports": [], "runs": []})
        resp = client.post("/api/repair-reports/analyze")
    assert resp.status_code == 400


@pytest.mark.asyncio
async def test_start_analysis_rejects_concurrent_run(client):
    active = _live_run(
        created=_iso_now_minus(1),
        started_at=_iso_now_minus(1),
    )
    with patch.object(
        reports, "repo_query", new_callable=AsyncMock
    ) as mock_repo:
        mock_repo.side_effect = _repo_router(
            {
                "reports": [_report_row()],
                "runs": [active],
                "commands": [
                    {
                        "status": "running",
                        "updated_at": _iso_now_minus(1),
                    }
                ],
            }
        )
        resp = client.post("/api/repair-reports/analyze")
    assert resp.status_code == 409


@pytest.mark.asyncio
async def test_start_analysis_submits_command_and_queues(client):
    updates = []

    async def _repo(query, params=None):
        if "FROM repair_analysis_run" in query:
            return []
        if "FROM repair_report ORDER BY" in query:
            return [_report_row(), _report_row(id="repair_report:def456")]
        if query.startswith("CREATE repair_analysis_run"):
            return [_run_row()]
        if query.startswith("UPDATE"):
            updates.append((query, params))
            return []
        raise AssertionError(f"unexpected query: {query}")

    with (
        patch.object(reports, "repo_query", new_callable=AsyncMock) as mock_repo,
        patch(
            "api.command_service.CommandService.submit_command_job",
            new=AsyncMock(return_value="command:job1"),
        ),
    ):
        mock_repo.side_effect = _repo
        resp = client.post("/api/repair-reports/analyze")
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["run"]["command_id"] == "command:job1"
    assert body["run"]["report_ids"] == [
        "repair_report:abc123",
        "repair_report:def456",
    ]
    queued = [
        params
        for query, params in updates
        if "analysis_state" in query and params.get("state") == "queued"
    ]
    assert len(queued) == 2


@pytest.mark.asyncio
async def test_stale_run_self_heals_and_unblocks():
    """A run whose worker died is finalized; the next analyze may proceed."""
    stale = _run_row(status="processing", command_id="command:dead")
    seen_updates = []

    async def _repo(query, params=None):
        if "FROM repair_analysis_run" in query:
            return [stale]
        if "FROM command" in query:
            return [{"status": "failed"}]
        if query.startswith("UPDATE"):
            seen_updates.append(query)
            return []
        if "FROM repair_report WHERE" in query:
            return [_report_row(analysis_state="processing")]
        raise AssertionError(f"unexpected query: {query}")

    with patch.object(reports, "repo_query", new_callable=AsyncMock) as mock_repo:
        mock_repo.side_effect = _repo
        assert await reports.get_active_run() is None
    assert any("finished_at" in query for query in seen_updates)


# --- SurrealQL statement integrity -----------------------------------------------
#
# Mocked persistence cannot catch brace/escaping mistakes in query strings
# (a doubled ``}}`` in a plain string is a permanent live failure). These
# tests pin the exact rendered statements.


@pytest.mark.asyncio
async def test_create_run_statement_is_valid_surrealql():
    seen = []

    async def _repo(query, params=None):
        seen.append(query)
        return [_run_row()]

    with patch.object(reports, "repo_query", new_callable=AsyncMock) as mock_repo:
        mock_repo.side_effect = _repo
        await reports.create_run(["repair_report:abc123"])
    (statement,) = seen
    assert statement == (
        "CREATE repair_analysis_run CONTENT {report_ids: $report_ids, "
        "manifest: [], status: $status, command_id: NONE, error: NONE, "
        "record_count: NONE, equipment_count: NONE, failure_mode_count: NONE, "
        "guide_count: NONE, created: time::now(), started_at: NONE, "
        "finished_at: NONE} RETURN AFTER"
    )


@pytest.mark.asyncio
async def test_create_report_statement_is_valid_surrealql(tmp_path, monkeypatch):
    monkeypatch.setattr(reports, "REPAIR_REPORTS_FOLDER", str(tmp_path))
    seen = []

    async def _repo(query, params=None):
        seen.append(query)
        return [_report_row()]

    content = _workbook_bytes(_sample_rows(count=1))
    with patch.object(reports, "repo_query", new_callable=AsyncMock) as mock_repo:
        mock_repo.side_effect = _repo
        await reports.create_report("cmms.xlsx", content)
    (statement,) = seen
    assert statement == (
        "CREATE repair_report CONTENT {"
        "filename: $filename, stored_filename: $stored_filename, "
        "size_bytes: $size_bytes, sheet: $sheet, column_count: $column_count, "
        "data_rows: $data_rows, analysis_key: $analysis_key, "
        "analysis_state: $analysis_state, last_run_id: NONE, "
        "last_completed_run_id: NONE, last_error: NONE, "
        "created: time::now(), updated: time::now()} RETURN AFTER"
    )


# --- aggregate builder (pure, no DB) --------------------------------------------


def test_storage_unique_naming_and_traversal_guard(tmp_path, monkeypatch):
    monkeypatch.setattr(reports, "REPAIR_REPORTS_FOLDER", str(tmp_path))
    first = reports.generate_unique_filename("cmms.xlsx")
    (open(first, "wb")).close()  # occupy the claimed name
    second = reports.generate_unique_filename("cmms.xlsx")
    assert first != second
    assert Path(second).parent == tmp_path
    # Directory components are stripped, never honored.
    sneaky = reports.generate_unique_filename("../evil.xlsx")
    assert Path(sneaky).parent == tmp_path
    assert Path(sneaky).name == "evil.xlsx"
    with pytest.raises(ValueError):
        reports.generate_unique_filename("")


def test_aggregate_namespaces_record_ids(tmp_path):
    file_a = tmp_path / "a.xlsx"
    file_b = tmp_path / "b.xlsx"
    file_a.write_bytes(_workbook_bytes(_sample_rows(prefix="B", start=1, count=2)))
    # Same natural IDs as file A: without namespacing these would collide.
    file_b.write_bytes(_workbook_bytes(_sample_rows(prefix="B", start=1, count=2)))
    dest = tmp_path / "aggregate.xlsx"
    manifest = reports.build_aggregate_workbook(
        [
            {
                "report_id": "repair_report:aaa",
                "filename": "a.xlsx",
                "analysis_key": "11111111",
                "path": str(file_a),
            },
            {
                "report_id": "repair_report:bbb",
                "filename": "b.xlsx",
                "analysis_key": "22222222",
                "path": str(file_b),
            },
        ],
        dest,
    )
    assert [entry["report_id"] for entry in manifest] == [
        "repair_report:aaa",
        "repair_report:bbb",
    ]
    assert manifest[0]["row_count"] == 2
    assert manifest[0]["first_row"] == 2
    assert manifest[1]["first_row"] == 4

    import openpyxl

    book = openpyxl.load_workbook(str(dest), read_only=True, data_only=True)
    sheet = book.active
    rows = list(sheet.iter_rows(values_only=True))
    headers = list(rows[0])
    prefix_pos = headers.index("پیشوند درخواست")
    prefixes = [row[prefix_pos] for row in rows[1:]]
    assert prefixes == ["11111111-B", "11111111-B", "22222222-B", "22222222-B"]
    book.close()


def test_namespaced_prefix_preserves_empty_cells():
    assert reports._namespaced_prefix("aa", None) is None
    assert reports._namespaced_prefix("aa", "  ") is None
    assert reports._namespaced_prefix("aa", "BR") == "aa-BR"
    assert reports._namespaced_prefix("aa", 210) == "aa-210"


# --- lifecycle consistency (Milestone 8) ------------------------------------------
#
# One truth across report state, run state, and command state: an active
# analysis never presents as completed, duplicates are refused with 409
# even inside the submit window, and dead workers heal exactly once.


@pytest.mark.asyncio
async def test_second_start_inside_submit_window_conflicts(client):
    """A command-less run younger than the grace window blocks duplicates."""
    young = _run_row(status="queued", command_id=None, created=_iso_now_minus(1))
    with patch.object(
        reports, "repo_query", new_callable=AsyncMock
    ) as mock_repo:
        mock_repo.side_effect = _repo_router(
            {"reports": [_report_row()], "runs": [young]}
        )
        resp = client.post("/api/repair-reports/analyze")
    assert resp.status_code == 409


@pytest.mark.asyncio
async def test_old_commandless_run_heals_then_allows_start(client):
    """A command-less run past the grace window heals; the next start works."""
    stale = _run_row(
        status="queued", command_id=None, created=_iso_now_minus(60)
    )
    updates = []

    async def _repo(query, params=None):
        if "FROM repair_analysis_run" in query:
            return [stale]
        if "FROM repair_report ORDER BY" in query:
            return [_report_row()]
        if "FROM repair_report WHERE" in query:
            return [_report_row()]
        if query.startswith("CREATE repair_analysis_run"):
            return [_run_row(id="repair_analysis_run:fresh")]
        if query.startswith("UPDATE"):
            updates.append((query, params))
            return []
        raise AssertionError(f"unexpected query: {query}")

    with (
        patch.object(reports, "repo_query", new_callable=AsyncMock) as mock_repo,
        patch(
            "api.command_service.CommandService.submit_command_job",
            new=AsyncMock(return_value="command:fresh"),
        ),
    ):
        mock_repo.side_effect = _repo
        resp = client.post("/api/repair-reports/analyze")
    assert resp.status_code == 200, resp.text
    assert resp.json()["run"]["id"] == "repair_analysis_run:fresh"
    assert any("finished_at" in query for query, _ in updates)


@pytest.mark.asyncio
async def test_new_command_blocks_start(client):
    """A `new` command recovers on worker restart, so it stays active."""
    active = _live_run(created=_iso_now_minus(1))
    with patch.object(
        reports, "repo_query", new_callable=AsyncMock
    ) as mock_repo:
        mock_repo.side_effect = _repo_router(
            {
                "reports": [_report_row()],
                "runs": [active],
                "commands": [{"status": "new"}],
            }
        )
        resp = client.post("/api/repair-reports/analyze")
    assert resp.status_code == 409


@pytest.mark.asyncio
async def test_orphaned_running_command_heals(client):
    """A `running` command without a fresh heartbeat is a dead worker."""
    orphan = _live_run(
        created=_iso_now_minus(60), started_at=_iso_now_minus(60)
    )
    seen_updates = []

    async def _repo(query, params=None):
        if "FROM repair_analysis_run" in query:
            return [orphan]
        if "FROM command" in query:
            return [{"status": "running"}]
        if "FROM repair_report WHERE" in query:
            return [_report_row(analysis_state="processing")]
        if query.startswith("UPDATE"):
            seen_updates.append(query)
            return []
        raise AssertionError(f"unexpected query: {query}")

    with patch.object(reports, "repo_query", new_callable=AsyncMock) as mock_repo:
        mock_repo.side_effect = _repo
        assert await reports.get_active_run() is None
    assert any("finished_at" in query for query in seen_updates)


@pytest.mark.asyncio
async def test_attach_failure_marks_run_failed():
    """Submit/attach fallout never leaves an orphan command-less run."""
    updates = []

    async def _repo(query, params=None):
        if "FROM repair_analysis_run" in query:
            return []
        if "FROM repair_report ORDER BY" in query:
            return [_report_row()]
        if query.startswith("CREATE repair_analysis_run"):
            return [_run_row()]
        if query.startswith("UPDATE"):
            updates.append((query, params))
            return []
        raise AssertionError(f"unexpected query: {query}")

    with (
        patch.object(reports, "repo_query", new_callable=AsyncMock) as mock_repo,
        patch(
            "api.command_service.CommandService.submit_command_job",
            new=AsyncMock(return_value="command:orphan"),
        ),
        patch.object(
            reports,
            "attach_command",
            new=AsyncMock(side_effect=RuntimeError("db down")),
        ),
    ):
        mock_repo.side_effect = _repo
        with pytest.raises(RuntimeError):
            await reports.start_analysis()
    failed = [
        params
        for query, params in updates
        if "repair_analysis_run" in str(params.get("rid", ""))
        and params.get("status") == "failed"
    ]
    assert failed, "expected the run to be marked failed after attach fallout"


@pytest.mark.asyncio
async def test_heartbeat_once_writes_liveness():
    from commands import repair_report_commands as worker_commands

    with patch(
        "open_notebook.database.repository.repo_query", new_callable=AsyncMock
    ) as mock_repo:
        mock_repo.return_value = []
        await worker_commands._heartbeat_once("command:abc")
    (query, params), _ = mock_repo.call_args
    assert "analysis_heartbeat" in query
    assert "updated_at" in query


# --- stale command flip + completion guard (Milestone 8) --------------------------


@pytest.mark.asyncio
async def test_finalize_flips_lease_expired_command():
    """A dead worker's `running` command is failed, not left running."""
    orphan = _live_run(
        created=_iso_now_minus(60), started_at=_iso_now_minus(60)
    )
    seen = []

    async def _repo(query, params=None):
        if "FROM repair_analysis_run" in query:
            return [orphan]
        if "FROM command" in query:
            return [{"id": "command:live", "status": "running"}]
        if "FROM repair_report WHERE" in query:
            return []
        if query.startswith("UPDATE"):
            seen.append((query, params))
            return []
        raise AssertionError(f"unexpected query: {query}")

    with patch.object(reports, "repo_query", new_callable=AsyncMock) as mock_repo:
        mock_repo.side_effect = _repo
        assert await reports.get_active_run() is None
    flipped = [
        params
        for query, params in seen
        if "error_message" in query and params.get("status") == "failed"
    ]
    assert flipped, "expected the orphaned command to be marked failed"


@pytest.mark.asyncio
async def test_finalize_keeps_new_command_for_restart():
    """A `new` command is never flipped: restarted workers resume it."""
    pending = _live_run(created=_iso_now_minus(60))
    seen = []

    async def _repo(query, params=None):
        if "FROM repair_analysis_run" in query:
            return [pending]
        if "FROM command" in query:
            return [{"id": "command:live", "status": "new"}]
        if query.startswith("UPDATE"):
            seen.append(query)
            return []
        raise AssertionError(f"unexpected query: {query}")

    with patch.object(reports, "repo_query", new_callable=AsyncMock) as mock_repo:
        mock_repo.side_effect = _repo
        active = await reports.get_active_run()
    assert active is not None
    assert not seen, "a `new` command must survive for worker restart"


@pytest.mark.asyncio
async def test_worker_refuses_to_complete_failed_run():
    """A superseded run can never overwrite the previous valid database."""
    from types import SimpleNamespace

    from commands import repair_report_commands as worker_commands

    completed_calls = []

    async def _fake_get_run(run_id):
        # Entry check: processing. Completion guard re-read: failed.
        if not completed_calls:
            completed_calls.append(1)
            return _live_run(status="processing")
        return _live_run(status="failed", error="healed while running")

    fake_result = SimpleNamespace(records=[], equipment=[], failure_modes=[], guides=[])
    fake_db = SimpleNamespace(parent=SimpleNamespace(mkdir=lambda **kwargs: None))

    with (
        patch.object(
            worker_commands.reports, "_get_run_internal", new=_fake_get_run
        ),
        patch.object(
            worker_commands.reports, "get_active_run", new=AsyncMock(return_value=None)
        ),
        patch.object(
            worker_commands.reports,
            "_get_report_internal",
            new=AsyncMock(return_value=_report_row()),
        ),
        patch.object(
            worker_commands.reports,
            "stored_path",
            return_value=SimpleNamespace(exists=lambda: True),
        ),
        patch.object(
            worker_commands.reports, "_set_report_state", new=AsyncMock()
        ),
        patch.object(
            worker_commands.reports, "mark_run_processing", new=AsyncMock()
        ),
        patch.object(
            worker_commands.reports, "mark_run_completed", new=AsyncMock()
        ) as mock_completed,
        patch.object(
            worker_commands.reports, "mark_run_failed", new=AsyncMock()
        ) as mock_failed,
        patch.object(
            worker_commands,
            "_build_aggregate",
            new=AsyncMock(return_value=[]),
        ),
        patch.object(
            worker_commands,
            "_load_engine",
            return_value={
                "analyze_workbook": lambda *args, **kwargs: fake_result,
                "EngineConfig": SimpleNamespace(default=lambda: None),
            },
        ),
        patch.object(
            worker_commands, "_resolve_database_path", return_value=fake_db
        ),
        patch(
            "open_notebook.database.repository.repo_query",
            new=AsyncMock(return_value=[]),
        ),
    ):
        with pytest.raises(ValueError, match="refusing to complete"):
            await worker_commands.analyze_repair_reports_command(
                worker_commands.AnalyzeRepairReportsInput(run_id="repair_analysis_run:x")
            )
    mock_completed.assert_not_called()
    mock_failed.assert_called_once()


# --- tasks endpoint: repair-analysis rows (Milestone 8) ----------------------------


def _tasks_repo_factory(embed_commands, analyze_commands, runs, reports):
    async def _repo(query, params=None):
        if "name = 'embed_source'" in query:
            return embed_commands
        if "analyze_repair_reports" in query:
            return analyze_commands
        if "FROM repair_analysis_run" in query:
            return runs
        if "FROM repair_report" in query:
            return reports
        if "FROM source" in query:
            return []
        raise AssertionError(f"unexpected query: {query}")

    return _repo


def _analyze_cmd(**overrides):
    record = {
        "id": "command:analysis1",
        "status": "running",
        "args": {"run_id": "repair_analysis_run:run1"},
        "result": None,
        "error_message": None,
        "created": "2026-09-27T00:00:00",
        "updated": None,
        "started_at": None,
        "updated_at": None,
    }
    record.update(overrides)
    return record


@pytest.mark.asyncio
@patch("open_notebook.database.repository.repo_query", new_callable=AsyncMock)
async def test_tasks_lists_active_repair_analysis(mock_repo, client):
    mock_repo.side_effect = _tasks_repo_factory(
        [],
        [_analyze_cmd()],
        [
            {
                "id": "repair_analysis_run:run1",
                "report_ids": [
                    "repair_report:aaa",
                    "repair_report:bbb",
                ],
            }
        ],
        [
            {"id": "repair_report:aaa", "filename": "a.xlsx"},
            {"id": "repair_report:bbb", "filename": "b.xlsx"},
        ],
    )
    (task,) = client.get("/api/tasks").json()
    assert task["item_type"] == "repair_analysis"
    assert task["command_name"] == "analyze_repair_reports"
    assert task["run_id"] == "repair_analysis_run:run1"
    assert task["title"] == "a.xlsx, b.xlsx"
    assert task["status"] == "running"
    # No chunk counts exist for analysis jobs: indeterminate, never faked.
    assert task["percentage"] is None
    assert task["processed_chunks"] is None


@pytest.mark.asyncio
@patch("open_notebook.database.repository.repo_query", new_callable=AsyncMock)
async def test_tasks_completed_repair_is_terminal(mock_repo, client):
    mock_repo.side_effect = _tasks_repo_factory(
        [],
        [_analyze_cmd(status="completed")],
        [{"id": "repair_analysis_run:run1", "report_ids": []}],
        [],
    )
    (task,) = client.get("/api/tasks").json()
    assert task["status"] == "completed"
    assert task["percentage"] == 100.0


@pytest.mark.asyncio
@patch("open_notebook.database.repository.repo_query", new_callable=AsyncMock)
async def test_tasks_failed_repair_carries_error(mock_repo, client):
    mock_repo.side_effect = _tasks_repo_factory(
        [],
        [_analyze_cmd(status="failed", error_message="boom")],
        [{"id": "repair_analysis_run:run1", "report_ids": []}],
        [],
    )
    (task,) = client.get("/api/tasks").json()
    assert task["status"] == "failed"
    assert task["error_message"] == "boom"


# --- end-to-end chain (real engine, tmp only) ------------------------------------


def test_chain_aggregate_engine_runtime(tmp_path, monkeypatch):
    """upload bytes → aggregate → engine → runtime adapter reads."""
    from maintenance_troubleshooting import EngineConfig, analyze_workbook

    import api.troubleshooting_service as troubleshooting_service

    file_a = tmp_path / "a.xlsx"
    file_b = tmp_path / "b.xlsx"
    file_a.write_bytes(_workbook_bytes(_sample_rows(prefix="B", start=1, count=3)))
    file_b.write_bytes(_workbook_bytes(_sample_rows(prefix="M", start=1, count=3)))

    dest = tmp_path / "aggregate.xlsx"
    reports.build_aggregate_workbook(
        [
            {
                "report_id": "repair_report:aaa",
                "filename": "a.xlsx",
                "analysis_key": "11111111",
                "path": str(file_a),
            },
            {
                "report_id": "repair_report:bbb",
                "filename": "b.xlsx",
                "analysis_key": "22222222",
                "path": str(file_b),
            },
        ],
        dest,
    )
    db_path = tmp_path / "knowledge.db"
    result = analyze_workbook(
        dest, configuration=EngineConfig.default(), output_path=db_path
    )
    assert len(result.records) == 6
    record_ids = [record.record_id for record in result.records]
    assert len(set(record_ids)) == 6
    assert any(rid.startswith("11111111-") for rid in record_ids)
    assert any(rid.startswith("22222222-") for rid in record_ids)

    monkeypatch.setenv("TROUBLESHOOTING_DB_PATH", str(db_path))
    status = troubleshooting_service.get_status()
    assert status["state"] == "available"
    equipment = troubleshooting_service.list_equipment()
    assert {item["code"] for item in equipment} == {"B104"}
    modes = troubleshooting_service.list_failure_modes("B104")
    assert len(modes) >= 1
    guide = troubleshooting_service.get_guide("B104", modes[0]["id"])
    assert guide["causes"], "expected mined causes"
    evidence_ids = [
        row["record_id"]
        for cause in guide["causes"]
        for row in cause["evidence"]
    ]
    assert evidence_ids, "expected provenance on causes"
    assert any(
        rid.startswith(("11111111-", "22222222-")) for rid in evidence_ids
    ), "evidence must trace to namespaced source records"


# --- single-report analysis + same-filename isolation (M11C-6R Part C) -----------


@pytest.mark.asyncio
async def test_start_single_report_analysis_queues_only_requested(client):
    """Per-report POST snapshots one file; no implicit process-everything."""
    updates = []

    async def _repo(query, params=None):
        if "FROM repair_analysis_run" in query:
            return []
        if "FROM repair_report WHERE" in query:
            return [_report_row(id="repair_report:abc123")]
        if query.startswith("CREATE repair_analysis_run"):
            return [_run_row(report_ids=["repair_report:abc123"])]
        if query.startswith("UPDATE"):
            updates.append((query, params))
            return []
        raise AssertionError(f"unexpected query: {query}")

    with (
        patch.object(reports, "repo_query", new_callable=AsyncMock) as mock_repo,
        patch(
            "api.command_service.CommandService.submit_command_job",
            new=AsyncMock(return_value="command:job1"),
        ),
    ):
        mock_repo.side_effect = _repo
        resp = client.post("/api/repair-reports/repair_report:abc123/analyze")
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["run"]["report_ids"] == ["repair_report:abc123"]
    queued = [
        params
        for query, params in updates
        if "analysis_state" in query and params.get("state") == "queued"
    ]
    assert len(queued) == 1


@pytest.mark.asyncio
async def test_start_single_report_unknown_is_404(client):
    async def _repo(query, params=None):
        if "FROM repair_analysis_run" in query:
            return []
        if "FROM repair_report WHERE" in query:
            return []
        raise AssertionError(f"unexpected query: {query}")

    with patch.object(reports, "repo_query", new_callable=AsyncMock) as mock_repo:
        mock_repo.side_effect = _repo
        resp = client.post("/api/repair-reports/repair_report:missing/analyze")
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_start_single_report_completed_refuses_reprocess(client):
    async def _repo(query, params=None):
        if "FROM repair_analysis_run" in query:
            return []
        if "FROM repair_report WHERE" in query:
            return [_report_row(analysis_state="completed")]
        raise AssertionError(f"unexpected query: {query}")

    with patch.object(reports, "repo_query", new_callable=AsyncMock) as mock_repo:
        mock_repo.side_effect = _repo
        resp = client.post("/api/repair-reports/repair_report:abc123/analyze")
    assert resp.status_code == 400


def test_same_filename_uploads_stay_isolated(tmp_path, monkeypatch):
    """Two uploads named cmms.xlsx never share storage or analysis keys."""
    monkeypatch.setattr(reports, "REPAIR_REPORTS_FOLDER", str(tmp_path))
    first = reports.generate_unique_filename("cmms.xlsx")
    Path(first).write_bytes(b"x")
    second = reports.generate_unique_filename("cmms.xlsx")
    assert first != second
    assert Path(first).name != Path(second).name
    # Analysis keys are per-file random hex, never derived from filenames.
    assert reports._namespaced_prefix("key-one", "B") != reports._namespaced_prefix(
        "key-two", "B"
    )


# --- report-scoped actions (M11C-6R Part B) -------------------------------------


def test_report_actions_filters_by_analysis_key(tmp_path, monkeypatch):
    """Actions/verifications/events split cleanly per analysis_key prefix."""

    from maintenance_troubleshooting import EngineConfig, analyze_workbook

    file_a = tmp_path / "a.xlsx"
    file_b = tmp_path / "b.xlsx"
    file_a.write_bytes(_workbook_bytes(_sample_rows(prefix="B", start=1, count=2)))
    file_b.write_bytes(_workbook_bytes(_sample_rows(prefix="M", start=1, count=2)))
    dest = tmp_path / "aggregate.xlsx"
    manifest = reports.build_aggregate_workbook(
        [
            {
                "report_id": "repair_report:aaa",
                "filename": "a.xlsx",
                "analysis_key": "11111111",
                "path": str(file_a),
            },
            {
                "report_id": "repair_report:bbb",
                "filename": "b.xlsx",
                "analysis_key": "22222222",
                "path": str(file_b),
            },
        ],
        dest,
    )
    db_path = tmp_path / "knowledge.db"
    analyze_workbook(dest, configuration=EngineConfig.default(), output_path=db_path)

    completed_run = _run_row(
        id="repair_analysis_run:done",
        status="completed",
        report_ids=["repair_report:aaa", "repair_report:bbb"],
        manifest=manifest,
    )

    async def _repo(query, params=None):
        if "FROM repair_report WHERE" in query:
            rid = str((params or {}).get("rid", ""))
            if "aaa" in rid:
                return [_report_row(id="repair_report:aaa", analysis_key="11111111")]
            return [_report_row(id="repair_report:bbb", analysis_key="22222222")]
        if "status = " in query and "repair_analysis_run" in query:
            return [completed_run]
        raise AssertionError(f"unexpected query: {query}")

    async def _run():
        with (
            patch.object(reports, "repo_query", new_callable=AsyncMock) as mock_repo,
            patch(
                "api.troubleshooting_service.resolve_database_path",
                return_value=db_path,
            ),
        ):
            mock_repo.side_effect = _repo
            actions_a = await reports.get_report_actions("repair_report:aaa")
            actions_b = await reports.get_report_actions("repair_report:bbb")
        return actions_a, actions_b

    import asyncio

    actions_a, actions_b = asyncio.run(_run())
    assert actions_a["analysis_key"] == "11111111"
    assert actions_b["analysis_key"] == "22222222"
    assert actions_a["run_id"] == "repair_analysis_run:done"
    # Record sets are disjoint and namespaced; nothing invented.
    assert actions_a["record_ids"]
    assert actions_b["record_ids"]
    assert not (set(actions_a["record_ids"]) & set(actions_b["record_ids"]))
    assert all(rid.startswith("11111111-") for rid in actions_a["record_ids"])
    assert all(rid.startswith("22222222-") for rid in actions_b["record_ids"])
    # Every returned object points at this report's records only.
    for payload in (actions_a, actions_b):
        allowed = set(payload["record_ids"])
        for action in payload["repair_actions"]:
            assert set(action["source_record_ids"]) <= allowed
        for item in payload["verifications"] + payload["post_repair_events"]:
            assert item["record_id"] in allowed
        for rid in payload["history_only_record_ids"]:
            assert rid in allowed


@pytest.mark.asyncio
async def test_report_actions_unknown_report_is_404(client):
    async def _repo(query, params=None):
        if "FROM repair_report WHERE" in query:
            return []
        raise AssertionError(f"unexpected query: {query}")

    with patch.object(reports, "repo_query", new_callable=AsyncMock) as mock_repo:
        mock_repo.side_effect = _repo
        resp = client.get("/api/repair-reports/repair_report:missing/actions")
    assert resp.status_code == 404
